from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

BASE_DIR = Path(__file__).resolve().parent.parent
SOURCE_IMAGE = BASE_DIR / "mapa_curricular.png"
OUTPUT_FILE = Path(__file__).resolve().parent / "mapa_curricular_editable.pptx"
ASSET_DIR = Path(__file__).resolve().parent / "_generated_assets"

DARK_BLUE = RGBColor(47, 82, 173)
TEXT_BLUE = RGBColor(45, 79, 165)
TEXT_DARK = RGBColor(55, 55, 55)
LIGHT_GRAY = RGBColor(242, 242, 242)
PANEL_GRAY = RGBColor(232, 232, 232)
WHITE = RGBColor(255, 255, 255)


def rel_x(slide_width, value):
    return Emu(int(slide_width * value))


def rel_y(slide_height, value):
    return Emu(int(slide_height * value))


def add_picture_full(slide, slide_width, slide_height, image_path):
    slide.shapes.add_picture(str(image_path), 0, 0, width=slide_width, height=slide_height)


def extract_icon(source_image, crop_box_rel, output_name, blue_only=False):
    width, height = source_image.size
    left = int(width * crop_box_rel[0])
    top = int(height * crop_box_rel[1])
    right = int(width * crop_box_rel[2])
    bottom = int(height * crop_box_rel[3])

    cropped = source_image.crop((left, top, right, bottom)).convert("RGBA")
    pixels = cropped.load()
    for y in range(cropped.height):
        for x in range(cropped.width):
            r, g, b, a = pixels[x, y]
            if blue_only:
                if not (b > 110 and b > r + 20 and b > g + 20):
                    pixels[x, y] = (255, 255, 255, 0)
            elif r > 228 and g > 228 and b > 228:
                pixels[x, y] = (255, 255, 255, 0)

    bbox = cropped.getbbox()
    if bbox:
        cropped = cropped.crop(bbox)

    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    output_path = ASSET_DIR / output_name
    cropped.save(output_path)
    return output_path


def extract_region(source_image, crop_box_rel, output_name):
    width, height = source_image.size
    left = int(width * crop_box_rel[0])
    top = int(height * crop_box_rel[1])
    right = int(width * crop_box_rel[2])
    bottom = int(height * crop_box_rel[3])

    cropped = source_image.crop((left, top, right, bottom)).convert("RGBA")
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    output_path = ASSET_DIR / output_name
    cropped.save(output_path)
    return output_path


def add_icon_picture(slide, image_path, left, top, width, height):
    slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)


def add_shape(slide, left, top, width, height, fill, line=None, radius=True):
    shape_type = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = fill if line is None else line
    if radius and hasattr(shape, "adjustments") and len(shape.adjustments) > 0:
        shape.adjustments[0] = 0.12
    return shape


def add_ellipse(slide, left, top, width, height, fill, line, line_width=2):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.OVAL, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line
    shape.line.width = Pt(line_width)
    return shape


def add_text(slide, left, top, width, height, text, font_size, color, bold=False, align=PP_ALIGN.LEFT, fill=None, margin=0.04):
    textbox = slide.shapes.add_textbox(left, top, width, height)
    if fill is not None:
        textbox.fill.solid()
        textbox.fill.fore_color.rgb = fill
        textbox.line.color.rgb = fill
    else:
        textbox.fill.background()
        textbox.line.fill.background()
    text_frame = textbox.text_frame
    text_frame.clear()
    text_frame.word_wrap = True
    text_frame.margin_left = Inches(margin)
    text_frame.margin_right = Inches(margin)
    text_frame.margin_top = Inches(margin)
    text_frame.margin_bottom = Inches(margin)
    text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    lines = text.split("\n")
    for index, line in enumerate(lines):
        paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Arial"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = color
    return textbox


def add_line(slide, left, top, width, color, weight=1.5):
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, left, top, width, Emu(int(Inches(0.01))))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.color.rgb = color
    line.line.width = Pt(weight)
    return line


def add_badge(slide, left, top, size, text):
    add_ellipse(slide, left, top, size, size, DARK_BLUE, WHITE, 1.5)
    add_text(slide, left, top, size, size, text, 9, WHITE, True, PP_ALIGN.CENTER, fill=None, margin=0.01)


def add_center_emblem(slide, slide_width, slide_height):
    icon_left = rel_x(slide_width, 0.468)
    icon_top = rel_y(slide_height, 0.302)
    icon_width = rel_x(slide_width, 0.04)
    icon_height = rel_y(slide_height, 0.065)

    add_shape(slide, icon_left, icon_top, icon_width, icon_height, WHITE, DARK_BLUE)
    add_shape(
        slide,
        icon_left + rel_x(slide_width, 0.008),
        icon_top + rel_y(slide_height, 0.012),
        rel_x(slide_width, 0.008),
        rel_y(slide_height, 0.04),
        WHITE,
        DARK_BLUE,
        radius=False,
    )
    add_shape(
        slide,
        icon_left + rel_x(slide_width, 0.02),
        icon_top + rel_y(slide_height, 0.012),
        rel_x(slide_width, 0.014),
        rel_y(slide_height, 0.018),
        WHITE,
        DARK_BLUE,
        radius=False,
    )
    add_line(slide, icon_left + rel_x(slide_width, 0.02), icon_top + rel_y(slide_height, 0.036), rel_x(slide_width, 0.014), DARK_BLUE, 1)
    add_line(slide, icon_left + rel_x(slide_width, 0.022), icon_top + rel_y(slide_height, 0.028), rel_x(slide_width, 0.01), DARK_BLUE, 1)


def add_dot_pattern(slide, start_x, start_y, cols, rows, gap_x, gap_y, size, color):
    for row in range(rows):
        for col in range(cols):
            left = start_x + (gap_x * col)
            top = start_y + (gap_y * row)
            add_ellipse(slide, left, top, size, size, color, color, 0.5)


def add_fan_lines(slide, start_x, start_y, width, count, gap_y, color):
    for index in range(count):
        add_line(slide, start_x + Emu(index * 1200), start_y + (gap_y * index), width - Emu(index * 1800), color, 0.8)


def build_presentation():
    if not SOURCE_IMAGE.exists():
        raise FileNotFoundError(f"No encontré la imagen fuente: {SOURCE_IMAGE}")

    asset_files = {
        "creditos": ASSET_DIR / "creditos.png",
        "semestres": ASSET_DIR / "semestres.png",
        "perfil_ingreso": ASSET_DIR / "perfil_ingreso.png",
        "perfil_egreso": ASSET_DIR / "perfil_egreso.png",
        "perfil_profesional": ASSET_DIR / "perfil_profesional.png",
        "genero": ASSET_DIR / "genero.png",
        "horas": ASSET_DIR / "horas.png",
        "ciencias_basicas": ASSET_DIR / "ciencias_basicas.png",
        "ciencias_sociales": ASSET_DIR / "ciencias_sociales.png",
        "ciencias_de_la_ingenieria": ASSET_DIR / "ciencias_de_la_ingenieria.png",
        "economicas": ASSET_DIR / "economicas.png",
        "ingenieria_aplicada": ASSET_DIR / "ingenieria_aplicada.png",
        "otras": ASSET_DIR / "otras.png",
    }
    missing_assets = [str(path) for path in asset_files.values() if not path.exists()]
    if missing_assets:
        raise FileNotFoundError(f"Faltan assets requeridos: {', '.join(missing_assets)}")

    with Image.open(SOURCE_IMAGE) as image:
        width_px, height_px = image.size
        source_art = image.copy()

    presentation = Presentation()
    presentation.slide_width = Inches(10)
    presentation.slide_height = Emu(int(presentation.slide_width * height_px / width_px))
    slide_width = presentation.slide_width
    slide_height = presentation.slide_height

    blank_layout = presentation.slide_layouts[6]

    reference_slide = presentation.slides.add_slide(blank_layout)
    add_picture_full(reference_slide, slide_width, slide_height, SOURCE_IMAGE)

    editable_slide = presentation.slides.add_slide(blank_layout)

    add_shape(editable_slide, 0, 0, slide_width, slide_height, WHITE, WHITE, radius=False)
    add_shape(
        editable_slide,
        rel_x(slide_width, 0.12),
        rel_y(slide_height, 0.135),
        rel_x(slide_width, 0.72),
        rel_y(slide_height, 0.79),
        WHITE,
        WHITE,
        radius=False,
    )
    add_shape(
        editable_slide,
        rel_x(slide_width, 0.12),
        rel_y(slide_height, 0.135),
        rel_x(slide_width, 0.72),
        rel_y(slide_height, 0.79),
        WHITE,
        PANEL_GRAY,
        radius=False,
    )
    add_dot_pattern(
        editable_slide,
        rel_x(slide_width, 0.13),
        rel_y(slide_height, 0.145),
        10,
        4,
        rel_x(slide_width, 0.018),
        rel_y(slide_height, 0.014),
        rel_x(slide_width, 0.008),
        RGBColor(237, 237, 237),
    )
    add_fan_lines(
        editable_slide,
        rel_x(slide_width, 0.68),
        rel_y(slide_height, 0.14),
        rel_x(slide_width, 0.12),
        8,
        rel_y(slide_height, 0.007),
        RGBColor(239, 239, 239),
    )
    add_shape(
        editable_slide,
        rel_x(slide_width, 0.16),
        rel_y(slide_height, 0.205),
        rel_x(slide_width, 0.22),
        rel_y(slide_height, 0.215),
        LIGHT_GRAY,
        LIGHT_GRAY,
    )
    add_shape(
        editable_slide,
        rel_x(slide_width, 0.61),
        rel_y(slide_height, 0.205),
        rel_x(slide_width, 0.22),
        rel_y(slide_height, 0.215),
        LIGHT_GRAY,
        LIGHT_GRAY,
    )
    add_ellipse(
        editable_slide,
        rel_x(slide_width, 0.402),
        rel_y(slide_height, 0.19),
        rel_x(slide_width, 0.165),
        rel_y(slide_height, 0.168),
        WHITE,
        DARK_BLUE,
        3,
    )
    add_ellipse(
        editable_slide,
        rel_x(slide_width, 0.412),
        rel_y(slide_height, 0.201),
        rel_x(slide_width, 0.145),
        rel_y(slide_height, 0.146),
        WHITE,
        PANEL_GRAY,
        1,
    )
    add_icon_picture(
        editable_slide,
        ASSET_DIR / "emblem.png",
        rel_x(slide_width, 0.357),
        rel_y(slide_height, 0.155),
        rel_x(slide_width, 0.255),
        rel_y(slide_height, 0.22),
    )

    add_text(
        editable_slide,
        rel_x(slide_width, 0.185),
        rel_y(slide_height, 0.178),
        rel_x(slide_width, 0.205),
        rel_y(slide_height, 0.036),
        "PERFIL GENERAL DE INGRESO",
        8.5,
        WHITE,
        True,
        PP_ALIGN.CENTER,
        DARK_BLUE,
    )
    add_text(
        editable_slide,
        rel_x(slide_width, 0.18),
        rel_y(slide_height, 0.25),
        rel_x(slide_width, 0.208),
        rel_y(slide_height, 0.145),
        "Conocimientos\n• Es conveniente cursar el área de ciencias físico-matemáticas en el bachillerato.\n• Contar con conocimientos generales de física y cómputo.\n• Comprensión básica de inglés, por lo menos a nivel de comprensión de textos.\n\nHabilidades\n• Disposición para el trabajo en equipo.\n• Capacidad de análisis y síntesis.\n• Adaptabilidad a situaciones nuevas.\n• Creatividad.",
        7.2,
        TEXT_DARK,
        False,
        PP_ALIGN.LEFT,
        fill=LIGHT_GRAY,
        margin=0.06,
    )

    add_text(
        editable_slide,
        rel_x(slide_width, 0.606),
        rel_y(slide_height, 0.178),
        rel_x(slide_width, 0.192),
        rel_y(slide_height, 0.036),
        "PERFIL GENERAL DE EGRESO",
        8.5,
        WHITE,
        True,
        PP_ALIGN.CENTER,
        DARK_BLUE,
    )
    add_text(
        editable_slide,
        rel_x(slide_width, 0.606),
        rel_y(slide_height, 0.232),
        rel_x(slide_width, 0.194),
        rel_y(slide_height, 0.196),
        "Conocimientos\n• Conocimiento para el modelado matemático de fenómenos físicos y de operación.\nHabilidades\n• Potencial para la creación de nuevas tecnologías.\n• Actitud emprendedora, directiva y de liderazgo.\n• Sensibilidad social y ética profesional.\n• Ser factor de cambio.\n• Actitud creativa e innovadora.\n• Habilidad de comunicación oral y escrita.\n• Capacidad para trabajar en entornos inter y multidisciplinarios.",
        6.7,
        TEXT_DARK,
        False,
        PP_ALIGN.LEFT,
        fill=LIGHT_GRAY,
        margin=0.05,
    )

    add_text(
        editable_slide,
        rel_x(slide_width, 0.395),
        rel_y(slide_height, 0.382),
        rel_x(slide_width, 0.215),
        rel_y(slide_height, 0.032),
        "OBJETIVO DEL PLAN DE ESTUDIOS",
        8.2,
        WHITE,
        True,
        PP_ALIGN.CENTER,
        DARK_BLUE,
    )
    add_text(
        editable_slide,
        rel_x(slide_width, 0.24),
        rel_y(slide_height, 0.41),
        rel_x(slide_width, 0.51),
        rel_y(slide_height, 0.055),
        "Formar profesionales en ingeniería capaces de identificar, desarrollar, proponer e integrar tecnologías para la mejora y desarrollo de productos, procesos y sistemas aeroespaciales. Asimismo, desarrollar profesionales con habilidades directivas, éticas, sociales y humanas.",
        7.6,
        TEXT_DARK,
        False,
        PP_ALIGN.CENTER,
        fill=PANEL_GRAY,
        margin=0.06,
    )
    add_line(editable_slide, rel_x(slide_width, 0.24), rel_y(slide_height, 0.468), rel_x(slide_width, 0.51), RGBColor(220, 220, 220), 0.8)

    add_shape(
        editable_slide,
        rel_x(slide_width, 0.145),
        rel_y(slide_height, 0.542),
        rel_x(slide_width, 0.14),
        rel_y(slide_height, 0.102),
        LIGHT_GRAY,
        LIGHT_GRAY,
    )
    add_shape(
        editable_slide,
        rel_x(slide_width, 0.145),
        rel_y(slide_height, 0.678),
        rel_x(slide_width, 0.14),
        rel_y(slide_height, 0.125),
        LIGHT_GRAY,
        LIGHT_GRAY,
    )
    add_shape(
        editable_slide,
        rel_x(slide_width, 0.63),
        rel_y(slide_height, 0.83),
        rel_x(slide_width, 0.19),
        rel_y(slide_height, 0.125),
        LIGHT_GRAY,
        LIGHT_GRAY,
    )

    add_text(
        editable_slide,
        rel_x(slide_width, 0.166),
        rel_y(slide_height, 0.50),
        rel_x(slide_width, 0.178),
        rel_y(slide_height, 0.05),
        "PERFIL ESPECÍFICO\nDE INGRESO",
        8.5,
        WHITE,
        True,
        PP_ALIGN.CENTER,
        DARK_BLUE,
    )
    add_icon_picture(
        editable_slide,
        asset_files["perfil_ingreso"],
        rel_x(slide_width, 0.138),
        rel_y(slide_height, 0.492),
        rel_x(slide_width, 0.06),
        rel_y(slide_height, 0.068),
    )
    add_text(
        editable_slide,
        rel_x(slide_width, 0.167),
        rel_y(slide_height, 0.55),
        rel_x(slide_width, 0.184),
        rel_y(slide_height, 0.094),
        "• Tener interés por el área de las tecnologías aeroespaciales.\n• Tener interés por el sector aeronáutico y espacial.",
        7.2,
        TEXT_DARK,
        False,
        PP_ALIGN.LEFT,
        fill=LIGHT_GRAY,
        margin=0.06,
    )

    add_text(
        editable_slide,
        rel_x(slide_width, 0.167),
        rel_y(slide_height, 0.635),
        rel_x(slide_width, 0.184),
        rel_y(slide_height, 0.052),
        "PERFIL ESPECÍFICO\nDE EGRESO",
        8.5,
        WHITE,
        True,
        PP_ALIGN.CENTER,
        DARK_BLUE,
    )
    add_icon_picture(
        editable_slide,
        asset_files["perfil_egreso"],
        rel_x(slide_width, 0.138),
        rel_y(slide_height, 0.627),
        rel_x(slide_width, 0.06),
        rel_y(slide_height, 0.068),
    )
    add_text(
        editable_slide,
        rel_x(slide_width, 0.167),
        rel_y(slide_height, 0.686),
        rel_x(slide_width, 0.185),
        rel_y(slide_height, 0.112),
        "• Formación de amplio espectro en ingeniería aeroespacial.\n• Habilidad para diseñar, construir, operar, dar mantenimiento, innovar, evaluar, modelar, simular e interpretar sistemas y tecnologías aeroespaciales.",
        6.9,
        TEXT_DARK,
        False,
        PP_ALIGN.LEFT,
        fill=LIGHT_GRAY,
        margin=0.06,
    )

    add_text(
        editable_slide,
        rel_x(slide_width, 0.43),
        rel_y(slide_height, 0.498),
        rel_x(slide_width, 0.18),
        rel_y(slide_height, 0.05),
        "450 créditos totales",
        16,
        TEXT_BLUE,
        True,
        PP_ALIGN.LEFT,
        fill=WHITE,
        margin=0.1,
    )
    add_icon_picture(
        editable_slide,
        asset_files["creditos"],
        rel_x(slide_width, 0.36),
        rel_y(slide_height, 0.482),
        rel_x(slide_width, 0.05),
        rel_y(slide_height, 0.068),
    )
    add_line(editable_slide, rel_x(slide_width, 0.60), rel_y(slide_height, 0.505), rel_x(slide_width, 0.0015), DARK_BLUE, 1)
    add_text(
        editable_slide,
        rel_x(slide_width, 0.665),
        rel_y(slide_height, 0.498),
        rel_x(slide_width, 0.14),
        rel_y(slide_height, 0.05),
        "Se cursa en\n10 semestres",
        10.5,
        TEXT_BLUE,
        True,
        PP_ALIGN.LEFT,
        fill=WHITE,
        margin=0.05,
    )
    add_icon_picture(
        editable_slide,
        asset_files["semestres"],
        rel_x(slide_width, 0.618),
        rel_y(slide_height, 0.483),
        rel_x(slide_width, 0.055),
        rel_y(slide_height, 0.065),
    )

    add_text(
        editable_slide,
        rel_x(slide_width, 0.378),
        rel_y(slide_height, 0.558),
        rel_x(slide_width, 0.17),
        rel_y(slide_height, 0.03),
        "Distribuidos en:",
        9,
        TEXT_DARK,
        False,
        PP_ALIGN.LEFT,
        fill=WHITE,
        margin=0.02,
    )
    add_line(editable_slide, rel_x(slide_width, 0.475), rel_y(slide_height, 0.575), rel_x(slide_width, 0.34), DARK_BLUE)

    categories = [
        (0.405, 0.595, "Ciencias Básicas\n128 créditos", asset_files["ciencias_basicas"]),
        (0.625, 0.595, "Ciencias Sociales y Humanidades\n28 créditos", asset_files["ciencias_sociales"]),
        (0.405, 0.655, "Ciencias de la Ingeniería\n140 créditos", asset_files["ciencias_de_la_ingenieria"]),
        (0.625, 0.655, "Ciencias Económico Administrativas\n30 créditos", asset_files["economicas"]),
        (0.405, 0.715, "Ingeniería Aplicada y diseño\n96 créditos", asset_files["ingenieria_aplicada"]),
        (0.625, 0.715, "Otras Asignaturas Convenientes\n28 créditos", asset_files["otras"]),
    ]

    for x, y, text, icon_path in categories:
        add_icon_picture(
            editable_slide,
            icon_path,
            rel_x(slide_width, x - 0.048),
            rel_y(slide_height, y - 0.003),
            rel_x(slide_width, 0.032),
            rel_y(slide_height, 0.045),
        )
        add_text(
            editable_slide,
            rel_x(slide_width, x),
            rel_y(slide_height, y),
            rel_x(slide_width, 0.17),
            rel_y(slide_height, 0.05),
            text,
            9.5,
            TEXT_BLUE,
            True,
            PP_ALIGN.LEFT,
            fill=WHITE,
            margin=0.03,
        )

    add_line(editable_slide, rel_x(slide_width, 0.37), rel_y(slide_height, 0.788), rel_x(slide_width, 0.44), DARK_BLUE)

    add_text(
        editable_slide,
        rel_x(slide_width, 0.26),
        rel_y(slide_height, 0.842),
        rel_x(slide_width, 0.33),
        rel_y(slide_height, 0.045),
        "3200 horas teóricas    800 horas prácticas",
        12.5,
        TEXT_BLUE,
        True,
        PP_ALIGN.CENTER,
        fill=WHITE,
        margin=0.02,
    )
    add_icon_picture(
        editable_slide,
        asset_files["horas"],
        rel_x(slide_width, 0.17),
        rel_y(slide_height, 0.818),
        rel_x(slide_width, 0.06),
        rel_y(slide_height, 0.095),
    )

    add_line(editable_slide, rel_x(slide_width, 0.17), rel_y(slide_height, 0.888), rel_x(slide_width, 0.64), DARK_BLUE)

    add_text(
        editable_slide,
        rel_x(slide_width, 0.255),
        rel_y(slide_height, 0.905),
        rel_x(slide_width, 0.33),
        rel_y(slide_height, 0.05),
        "Asignatura Igualdad de Género\nen Ingeniería como requisito de permanencia",
        10.5,
        TEXT_BLUE,
        True,
        PP_ALIGN.LEFT,
        fill=WHITE,
        margin=0.03,
    )
    add_icon_picture(
        editable_slide,
        asset_files["genero"],
        rel_x(slide_width, 0.155),
        rel_y(slide_height, 0.885),
        rel_x(slide_width, 0.052),
        rel_y(slide_height, 0.08),
    )

    add_text(
        editable_slide,
        rel_x(slide_width, 0.632),
        rel_y(slide_height, 0.77),
        rel_x(slide_width, 0.18),
        rel_y(slide_height, 0.06),
        "PERFIL\nPROFESIONAL",
        9.5,
        WHITE,
        True,
        PP_ALIGN.CENTER,
        DARK_BLUE,
    )
    add_icon_picture(
        editable_slide,
        asset_files["perfil_profesional"],
        rel_x(slide_width, 0.637),
        rel_y(slide_height, 0.762),
        rel_x(slide_width, 0.06),
        rel_y(slide_height, 0.068),
    )
    add_text(
        editable_slide,
        rel_x(slide_width, 0.63),
        rel_y(slide_height, 0.83),
        rel_x(slide_width, 0.19),
        rel_y(slide_height, 0.125),
        "• Sector aeronáutico: diseño y manufactura, sistemas de navegación, materiales y pruebas de propulsión.\n• Sector aeroespacial: diseño de misiones espaciales, pruebas de certificación, desarrollo de subsistemas satelitales y sistemas de lanzamiento.\n• Comunicación y gestión tecnológica.\n• Capacidad para laborar en el sector público y privado.",
        6.6,
        TEXT_DARK,
        False,
        PP_ALIGN.LEFT,
        fill=LIGHT_GRAY,
        margin=0.05,
    )

    presentation.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    output = build_presentation()
    print(output)
