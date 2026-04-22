"""
Export the English curricular map PPTX content to a Word document (.docx)
with icons inline and formatted text, preserving the visual structure.
"""
import io
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

SRC_PPTX = Path(__file__).resolve().parent / "mapa_curricular_english_v2.pptx"
DST_DOCX = Path(__file__).resolve().parent / "mapa_curricular_english.docx"

BLUE = RGBColor(47, 82, 173)
DARK = RGBColor(55, 55, 55)


def extract_images_and_text(pptx_path):
    """Extract all shapes with their positions, images, and text."""
    prs = Presentation(str(pptx_path))
    shapes_data = []
    for slide in prs.slides:
        for shape in slide.shapes:
            info = {
                "name": shape.name,
                "left": shape.left,
                "top": shape.top,
                "width": shape.width,
                "height": shape.height,
                "is_picture": shape.shape_type == MSO_SHAPE_TYPE.PICTURE,
                "image_blob": None,
                "text_lines": [],
            }
            if info["is_picture"]:
                info["image_blob"] = shape.image.blob
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    txt = para.text.strip()
                    if txt:
                        bold = any(r.font.bold for r in para.runs if r.font.bold)
                        size = None
                        for r in para.runs:
                            if r.font.size:
                                size = r.font.size
                                break
                        info["text_lines"].append({
                            "text": txt,
                            "bold": bold,
                            "size": size,
                        })
            shapes_data.append(info)
    return shapes_data


def build_word(shapes_data):
    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Arial"
    font.size = Pt(11)

    # ── Title ──
    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.add_run("Aerospace Engineering — Curricular Map")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = BLUE
    run.font.name = "Arial"

    doc.add_paragraph()  # spacer

    # ── Aerospace emblem image ──
    for s in shapes_data:
        if s["name"] == "Picture 56" and s["image_blob"]:
            img_stream = io.BytesIO(s["image_blob"])
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = para.add_run()
            run.add_picture(img_stream, width=Inches(2.0))
            break

    doc.add_paragraph()  # spacer

    # ── Credits and Semesters (side by side table) ──
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    # Credits cell
    cell_credits = table.cell(0, 0)
    cell_credits.width = Inches(3.2)
    p = cell_credits.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for s in shapes_data:
        if s["name"] == "Picture 74" and s["image_blob"]:
            run = p.add_run()
            run.add_picture(io.BytesIO(s["image_blob"]), width=Inches(0.5))
            break
    run = p.add_run("  450 total credits")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = BLUE
    run.font.name = "Arial"

    # Semesters cell
    cell_sem = table.cell(0, 1)
    cell_sem.width = Inches(3.2)
    p = cell_sem.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for s in shapes_data:
        if s["name"] == "Picture 77" and s["image_blob"]:
            run = p.add_run()
            run.add_picture(io.BytesIO(s["image_blob"]), width=Inches(0.5))
            break
    run = p.add_run("  Completed in 10 semesters")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = BLUE
    run.font.name = "Arial"

    doc.add_paragraph()  # spacer

    # ── "Distributed in:" header ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Distributed in:")
    run.font.size = Pt(12)
    run.font.color.rgb = DARK
    run.font.name = "Arial"

    # ── Horizontal line ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("─" * 70)
    run.font.color.rgb = BLUE
    run.font.size = Pt(8)
    run.font.name = "Arial"

    # ── Categories table (3 rows x 2 cols) ──
    categories = [
        ("Picture 80", "Basic Sciences\n128 credits"),
        ("Picture 82", "Social Sciences and Humanities\n28 credits"),
        ("Picture 84", "Engineering Sciences\n140 credits"),
        ("Picture 86", "Economic-Administrative Sciences\n30 credits"),
        ("Picture 88", "Applied Engineering and Design\n96 credits"),
        ("Picture 90", "Other Elective Courses\n28 credits"),
    ]

    shape_map = {s["name"]: s for s in shapes_data}

    cat_table = doc.add_table(rows=3, cols=2)
    cat_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cat_table.autofit = True

    for idx, (pic_name, label) in enumerate(categories):
        row = idx // 2
        col = idx % 2
        cell = cat_table.cell(row, col)
        cell.width = Inches(3.2)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

        s = shape_map.get(pic_name)
        if s and s["image_blob"]:
            run = p.add_run()
            run.add_picture(io.BytesIO(s["image_blob"]), width=Inches(0.4))

        lines = label.split("\n")
        run = p.add_run(f"  {lines[0]}")
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = BLUE
        run.font.name = "Arial"

        if len(lines) > 1:
            run = p.add_run(f"\n       {lines[1]}")
            run.bold = True
            run.font.size = Pt(10)
            run.font.color.rgb = BLUE
            run.font.name = "Arial"

    doc.add_paragraph()  # spacer

    # ── Horizontal line ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("─" * 70)
    run.font.color.rgb = BLUE
    run.font.size = Pt(8)
    run.font.name = "Arial"

    # ── Hours ──
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("3,200 theoretical hours")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = BLUE
    run.font.name = "Arial"
    run = p.add_run("        ")
    run = p.add_run("800 practical hours")
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = BLUE
    run.font.name = "Arial"

    # ── Remove table borders ──
    for tbl in [table, cat_table]:
        for row in tbl.rows:
            for cell in row.cells:
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                from docx.oxml.ns import qn
                from lxml import etree
                borders = etree.SubElement(tcPr, qn("w:tcBorders"))
                for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                    el = etree.SubElement(borders, qn(f"w:{edge}"))
                    el.set(qn("w:val"), "none")
                    el.set(qn("w:sz"), "0")
                    el.set(qn("w:space"), "0")
                    el.set(qn("w:color"), "auto")

    doc.save(str(DST_DOCX))
    print(f"Saved to: {DST_DOCX}")


def main():
    print("Extracting from PPTX...")
    shapes_data = extract_images_and_text(str(SRC_PPTX))
    print(f"  Found {len(shapes_data)} shapes")

    print("Building Word document...")
    build_word(shapes_data)


if __name__ == "__main__":
    main()
