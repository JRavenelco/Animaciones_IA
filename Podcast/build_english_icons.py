"""
Generate new unique blue line-art icons using PIL and produce the final
English PPTX with Arial fonts and replaced icons.
"""
import io
import math
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

SRC = Path(__file__).resolve().parent / "mapa_curricular_english.pptx"
DST = Path(__file__).resolve().parent / "mapa_curricular_english_v2.pptx"

BLUE = (47, 82, 173)
BG = (0, 0, 0, 0)  # transparent
ICON_SIZE = 512
LW = 18  # base line width


def _canvas():
    img = Image.new("RGBA", (ICON_SIZE, ICON_SIZE), BG)
    draw = ImageDraw.Draw(img)
    return img, draw


# ── 1. Credits icon: a medal / ribbon ──
def icon_credits():
    img, d = _canvas()
    cx, cy = 256, 200
    r = 100
    d.ellipse([cx - r, cy - r, cx + r, cy + r], outline=BLUE, width=LW)
    # star inside
    pts = []
    for i in range(5):
        angle = math.radians(-90 + i * 72)
        pts.append((cx + int(70 * math.cos(angle)), cy + int(70 * math.sin(angle))))
    star = []
    for i in range(5):
        star.append(pts[i])
        mid_angle = math.radians(-90 + i * 72 + 36)
        star.append((cx + int(30 * math.cos(mid_angle)), cy + int(30 * math.sin(mid_angle))))
    d.polygon(star, outline=BLUE, width=LW // 2)
    # ribbon tails
    d.line([cx - 50, cy + r, cx - 90, cy + r + 140], fill=BLUE, width=LW)
    d.line([cx - 50, cy + r, cx - 20, cy + r + 80], fill=BLUE, width=LW // 2)
    d.line([cx + 50, cy + r, cx + 90, cy + r + 140], fill=BLUE, width=LW)
    d.line([cx + 50, cy + r, cx + 20, cy + r + 80], fill=BLUE, width=LW // 2)
    return img


# ── 2. Semesters icon: calendar with grid ──
def icon_semesters():
    img, d = _canvas()
    x0, y0, x1, y1 = 80, 100, 432, 432
    d.rounded_rectangle([x0, y0, x1, y1], radius=24, outline=BLUE, width=LW)
    # top bar
    d.rectangle([x0, y0, x1, y0 + 70], outline=BLUE, fill=BLUE)
    # calendar rings
    for rx in [170, 256, 342]:
        d.line([rx, y0 - 30, rx, y0 + 30], fill=BLUE, width=LW)
    # grid lines
    for gy in range(210, 420, 55):
        d.line([x0 + 30, gy, x1 - 30, gy], fill=BLUE, width=LW // 3)
    for gx in range(160, 420, 70):
        d.line([gx, 195, gx, 410], fill=BLUE, width=LW // 3)
    return img


# ── 3. Basic Sciences: atom ──
def icon_basic_sciences():
    img, d = _canvas()
    cx, cy = 256, 256
    # nucleus
    d.ellipse([cx - 22, cy - 22, cx + 22, cy + 22], fill=BLUE)
    # three elliptical orbits
    for angle_deg in [0, 60, 120]:
        a = math.radians(angle_deg)
        pts = []
        for t in range(0, 360, 3):
            tr = math.radians(t)
            ex = 160 * math.cos(tr)
            ey = 55 * math.sin(tr)
            rx = ex * math.cos(a) - ey * math.sin(a)
            ry = ex * math.sin(a) + ey * math.cos(a)
            pts.append((cx + int(rx), cy + int(ry)))
        for i in range(len(pts) - 1):
            d.line([pts[i], pts[i + 1]], fill=BLUE, width=LW // 2)
        d.line([pts[-1], pts[0]], fill=BLUE, width=LW // 2)
    return img


# ── 4. Social Sciences: three people silhouettes ──
def icon_social_sciences():
    img, d = _canvas()
    # center person
    d.ellipse([226, 80, 286, 140], outline=BLUE, width=LW)
    d.arc([186, 150, 326, 290], 0, 180, fill=BLUE, width=LW)
    # left person (smaller, behind)
    d.ellipse([106, 120, 156, 170], outline=BLUE, width=LW - 4)
    d.arc([76, 180, 186, 300], 0, 180, fill=BLUE, width=LW - 4)
    # right person (smaller, behind)
    d.ellipse([356, 120, 406, 170], outline=BLUE, width=LW - 4)
    d.arc([326, 180, 436, 300], 0, 180, fill=BLUE, width=LW - 4)
    # connecting base line
    d.arc([100, 260, 412, 420], 0, 180, fill=BLUE, width=LW // 2)
    return img


# ── 5. Engineering Sciences: gear + wrench ──
def icon_engineering_sciences():
    img, d = _canvas()
    cx, cy = 230, 230
    # gear outer
    teeth = 8
    for i in range(teeth):
        a = math.radians(i * 360 / teeth)
        x1 = cx + int(120 * math.cos(a))
        y1 = cy + int(120 * math.sin(a))
        d.rectangle([x1 - 20, y1 - 20, x1 + 20, y1 + 20], fill=BLUE)
    d.ellipse([cx - 100, cy - 100, cx + 100, cy + 100], outline=BLUE, width=LW)
    d.ellipse([cx - 40, cy - 40, cx + 40, cy + 40], outline=BLUE, width=LW)
    # wrench handle (diagonal)
    d.line([330, 330, 440, 440], fill=BLUE, width=LW + 4)
    d.ellipse([300, 300, 370, 370], outline=BLUE, width=LW)
    return img


# ── 6. Economic-Admin: bar chart with trend arrow ──
def icon_economics():
    img, d = _canvas()
    # axes
    d.line([100, 400, 100, 80], fill=BLUE, width=LW)
    d.line([100, 400, 430, 400], fill=BLUE, width=LW)
    # bars
    bars = [(140, 320), (210, 250), (280, 180), (350, 130)]
    bw = 50
    for bx, by in bars:
        d.rectangle([bx, by, bx + bw, 400], outline=BLUE, fill=BLUE + (100,), width=LW // 2)
        d.rectangle([bx, by, bx + bw, 400], outline=BLUE, width=LW // 2)
    # trend arrow
    d.line([150, 340, 390, 120], fill=BLUE, width=LW // 2)
    # arrowhead
    d.polygon([(390, 120), (360, 140), (370, 110)], fill=BLUE)
    return img


# ── 7. Applied Engineering: lightbulb + gear ──
def icon_applied_eng():
    img, d = _canvas()
    cx, cy = 256, 200
    # bulb top
    d.ellipse([cx - 90, cy - 120, cx + 90, cy + 60], outline=BLUE, width=LW)
    # filament
    d.line([cx - 25, cy - 40, cx, cy - 70, cx + 25, cy - 40], fill=BLUE, width=LW // 2)
    # base
    d.rectangle([cx - 45, cy + 60, cx + 45, cy + 120], outline=BLUE, width=LW)
    d.line([cx - 45, cy + 80, cx + 45, cy + 80], fill=BLUE, width=LW // 3)
    d.line([cx - 45, cy + 100, cx + 45, cy + 100], fill=BLUE, width=LW // 3)
    # small gear bottom-right
    gcx, gcy = 370, 370
    gr = 55
    d.ellipse([gcx - gr, gcy - gr, gcx + gr, gcy + gr], outline=BLUE, width=LW - 4)
    d.ellipse([gcx - 20, gcy - 20, gcx + 20, gcy + 20], outline=BLUE, width=LW - 4)
    for i in range(6):
        a = math.radians(i * 60)
        tx = gcx + int((gr + 12) * math.cos(a))
        ty = gcy + int((gr + 12) * math.sin(a))
        d.rectangle([tx - 10, ty - 10, tx + 10, ty + 10], fill=BLUE)
    return img


# ── 8. Other Courses: open book with bookmark ──
def icon_other_courses():
    img, d = _canvas()
    # left page
    d.polygon([(256, 130), (80, 100), (80, 390), (256, 400)], outline=BLUE, width=LW)
    # right page
    d.polygon([(256, 130), (432, 100), (432, 390), (256, 400)], outline=BLUE, width=LW)
    # spine
    d.line([256, 130, 256, 400], fill=BLUE, width=LW)
    # left page lines
    for ly in range(190, 370, 40):
        d.line([120, ly, 230, ly + 10], fill=BLUE, width=LW // 4)
    # right page lines
    for ly in range(190, 370, 40):
        d.line([282, ly + 10, 400, ly], fill=BLUE, width=LW // 4)
    # bookmark
    d.polygon([(370, 100), (370, 220), (390, 195), (410, 220), (410, 100)], outline=BLUE, fill=BLUE + (80,), width=LW // 3)
    return img


ICON_GENERATORS = {
    "credits": icon_credits,
    "semesters": icon_semesters,
    "basic_sciences": icon_basic_sciences,
    "social_sciences": icon_social_sciences,
    "engineering_sciences": icon_engineering_sciences,
    "economics": icon_economics,
    "applied_engineering": icon_applied_eng,
    "other_courses": icon_other_courses,
}

# Map shape names to icon generator keys
SHAPE_TO_ICON = {
    "Picture 74": "credits",
    "Picture 77": "semesters",
    "Picture 80": "basic_sciences",
    "Picture 82": "social_sciences",
    "Picture 84": "engineering_sciences",
    "Picture 86": "economics",
    "Picture 88": "applied_engineering",
    "Picture 90": "other_courses",
}


def generate_icon_bytes(key):
    img = ICON_GENERATORS[key]()
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def replace_image_in_shape(shape, png_bytes):
    """Replace the image blob inside an existing picture shape."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    rel = shape.part.rels[shape._element.blipFill.blip.get(
        "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
    )]
    rel.target_part._blob = png_bytes


def set_all_fonts_arial(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    for run in para.runs:
                        run.font.name = "Arial"


def main():
    print("Loading presentation...")
    prs = Presentation(str(SRC))

    print("Setting all fonts to Arial...")
    set_all_fonts_arial(prs)

    print("Generating and replacing icons...")
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.name in SHAPE_TO_ICON:
                icon_key = SHAPE_TO_ICON[shape.name]
                png_data = generate_icon_bytes(icon_key)
                replace_image_in_shape(shape, png_data)
                print(f"  Replaced {shape.name} -> {icon_key}")

    prs.save(str(DST))
    print(f"\nSaved to: {DST}")


if __name__ == "__main__":
    main()
